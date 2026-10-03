
from docx import Document

# Create a new Word document
doc = Document()

# Add project title
doc.add_heading(
    "E-Commerce Sales Analysis and Advanced Data Visualization",
    0
)

# Add subtitle
doc.add_paragraph(
    "Week 2 Internship Project Report"
)

# Save document
doc.save("report_test.docx")

print("Word document created successfully!")